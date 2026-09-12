# -*- coding: utf-8 -*-
import re
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"G:\sistemas de informacion 2 angelica garzon\fashonstore parcial si2\tienda_ropa\docs\documentacion_fase1_fashionstore.md"
OUT = r"G:\sistemas de informacion 2 angelica garzon\fashonstore parcial si2\documentos\FashionStore_Presentacion1.docx"

INLINE_RE = re.compile(r'(\*\*.+?\*\*|`[^`]+?`)')

def add_inline_runs(paragraph, text):
    parts = INLINE_RE.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = paragraph.add_run(part[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        else:
            paragraph.add_run(part)

def set_cell_shading(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def add_hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'B0B0B0')
    pBdr.append(bottom)
    pPr.append(pBdr)

def parse_table(lines, i):
    rows = []
    while i < len(lines) and lines[i].strip().startswith('|'):
        row = lines[i].strip()
        if not re.match(r'^\|[\s:\-\|]+\|$', row):
            cells = [c.strip() for c in row.strip('|').split('|')]
            rows.append(cells)
        i += 1
    return rows, i

def build_table(doc, rows):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for cidx in range(ncols):
            text = row[cidx] if cidx < len(row) else ''
            para = cells[cidx].paragraphs[0]
            add_inline_runs(para, text)
            for run in para.runs:
                run.font.size = Pt(9)
            if ridx == 0:
                for run in para.runs:
                    run.bold = True
                set_cell_shading(cells[cidx], '1F2937')
                for run in para.runs:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    doc.add_paragraph()

def main():
    with open(SRC, encoding='utf-8') as f:
        lines = f.read().split('\n')

    doc = Document()

    # base style
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)

    for i in range(1, 5):
        h = doc.styles[f'Heading {i}']
        h.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    sec = doc.sections[0]
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(2)
    sec.bottom_margin = Cm(2)

    i = 0
    in_code = False
    code_buffer = []

    def flush_code():
        nonlocal code_buffer
        if code_buffer:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            shading_elm = OxmlElement('w:shd')
            shading_elm.set(qn('w:fill'), 'F2F2F0')
            p._p.get_or_add_pPr().append(shading_elm)
            run = p.add_run('\n'.join(code_buffer))
            run.font.name = 'Consolas'
            run.font.size = Pt(8.5)
            code_buffer = []

    n = len(lines)
    while i < n:
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith('```'):
            if in_code:
                flush_code()
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_buffer.append(raw)
            i += 1
            continue

        if stripped == '':
            i += 1
            continue

        if stripped == '---':
            add_hr(doc)
            i += 1
            continue

        if stripped.startswith('|'):
            rows, i = parse_table(lines, i)
            build_table(doc, rows)
            continue

        m = re.match(r'^(#{1,5})\s+(.*)$', stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            if level == 1:
                p = doc.add_heading(level=0)
                add_inline_runs(p, text)
            else:
                p = doc.add_heading(level=min(level, 4))
                add_inline_runs(p, text)
            i += 1
            continue

        m = re.match(r'^[-*]\s+(.*)$', stripped)
        if m:
            p = doc.add_paragraph(style='List Bullet')
            add_inline_runs(p, m.group(1))
            i += 1
            continue

        m = re.match(r'^\*\((.*)\)\*$', stripped)
        if m:
            p = doc.add_paragraph()
            run = p.add_run(m.group(1))
            run.italic = True
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x6B, 0x6B, 0x6B)
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline_runs(p, stripped)
        i += 1

    doc.save(OUT)
    print('OK ->', OUT)

if __name__ == '__main__':
    main()
